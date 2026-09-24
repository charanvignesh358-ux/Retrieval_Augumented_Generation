"""Enterprise RAG Document Assistant - FastAPI + ChromaDB + Groq (free)."""
import os, io, uuid, time, json, logging, hmac, secrets, re
from pathlib import Path
from typing import Optional, Dict, List

from dotenv import load_dotenv
from fastapi import FastAPI, UploadFile, File, HTTPException, Depends, Header
from itsdangerous import URLSafeTimedSerializer, BadSignature, SignatureExpired
from fastapi.responses import FileResponse
from fastapi.staticfiles import StaticFiles
from pydantic import BaseModel
import chromadb
from groq import Groq
from pypdf import PdfReader
import docx

try:
    from neo4j import GraphDatabase
except Exception:  # pragma: no cover
    GraphDatabase = None

BASE = Path(__file__).parent
load_dotenv(BASE / ".env")
load_dotenv(BASE.parent / ".env")
GROQ_MODEL = os.getenv("GROQ_MODEL", "qwen/qwen3.8-27b")
GROQ_MODEL_FALLBACKS = [
    "qwen/qwen3.8-27b",
    "openai/gpt-oss-20b",
    "openai/gpt-oss-120b",
]
TOP_K = int(os.getenv("TOP_K", 5))
CHUNK_WORDS = int(os.getenv("CHUNK_WORDS", 350))
CHUNK_OVERLAP = int(os.getenv("CHUNK_OVERLAP", 60))
AUDIT_LOG = BASE / "audit.log"
OCR_ENABLED = os.getenv("OCR_ENABLED", "true").lower() == "true"
TOKEN_HOURS = int(os.getenv("TOKEN_HOURS", 8))

# ---------- auth config: APP_USERS="admin:pass:admin,alice:pass:user" ----------
USERS: Dict[str, tuple] = {}
for _item in os.getenv("APP_USERS", "").split(","):
    _p = _item.strip().split(":")
    if len(_p) >= 2:
        USERS[_p[0]] = (_p[1], _p[2] if len(_p) > 2 else "user")
if not USERS or any("change_me" in v[0] for v in USERS.values()):
    raise SystemExit("Set real passwords in APP_USERS inside groq_assistant/.env (format user:password:role).")
SECRET = os.getenv("SECRET_KEY", "")
if not SECRET or "change_me" in SECRET:
    SECRET = secrets.token_hex(32)  # random per start; set SECRET_KEY in .env to stay logged in across restarts
signer = URLSafeTimedSerializer(SECRET)
FAILS: Dict[str, list] = {}

logging.basicConfig(level=logging.INFO)
if not os.getenv("GROQ_API_KEY"):
    raise SystemExit("Set GROQ_API_KEY in groq_assistant/.env (free key: https://console.groq.com/keys)")

groq = Groq(api_key=os.environ["GROQ_API_KEY"])
# Chroma's default embedder (ONNX all-MiniLM-L6-v2) runs locally: no paid embedding API needed.
chroma = chromadb.PersistentClient(path=str(BASE / "chroma_db"))
col = chroma.get_or_create_collection("docs", metadata={"hnsw:space": "cosine"})

SYSTEM_PROMPT = (
    "You are the Enterprise Document Assistant. Answer ONLY using the provided CONTEXT. "
    "Cite sources inline like [Source: file, page N]. If the answer is not in the context, "
    "reply exactly: \"I don't know based on the available documents.\" "
    "Never invent facts. Be professional and concise."
)

GRAPH_RELATION_PATTERNS = [
    r"\b([A-Z][A-Za-z0-9/._+-]*(?:\s+[A-Z][A-Za-z0-9/._+-]*)*)\s+(uses|connects to|depends on|integrates with|stores|retrieves|sends data to|calls|provides|contains|includes|reads|writes to|works with)\s+([A-Z][A-Za-z0-9/._+-]*(?:\s+[A-Z][A-Za-z0-9/._+-]*)*)",
    r"\b([A-Z][A-Za-z0-9/._+-]*(?:\s+[A-Z][A-Za-z0-9/._+-]*)*)\s+(is|are)\s+(?:a|an|the)?\s*([A-Z][A-Za-z0-9/._+-]*(?:\s+[A-Z][A-Za-z0-9/._+-]*)*)",
    r"\b([A-Z][A-Za-z0-9/._+-]*(?:\s+[A-Z][A-Za-z0-9/._+-]*)*)\s+(includes|contains|supports)\s+([A-Z][A-Za-z0-9/._+-]*(?:\s+[A-Z][A-Za-z0-9/._+-]*)*)",
]

_ocr_engine = None
GRAPH_INDEX: Dict[str, List[dict]] = {}
NEO4J_URI = os.getenv("NEO4J_URI")
NEO4J_USER = os.getenv("NEO4J_USER")
NEO4J_PASSWORD = os.getenv("NEO4J_PASSWORD")
NEO4J_DRIVER = None
if GraphDatabase and NEO4J_URI and NEO4J_USER and NEO4J_PASSWORD:
    try:
        NEO4J_DRIVER = GraphDatabase.driver(NEO4J_URI, auth=(NEO4J_USER, NEO4J_PASSWORD))
        with NEO4J_DRIVER.session() as session:
            session.run("RETURN 1")
    except Exception as exc:
        logging.warning("Neo4j not available: %s", exc)
        NEO4J_DRIVER = None


def normalize_entity(value: str) -> str:
    if not value:
        return ""
    value = re.sub(r"\s+", " ", value.strip())
    value = value.strip("()[]{}<>;,:.")
    value = value.strip("\"'")
    return value


def extract_graph_facts(text: str):
    facts: List[dict] = []
    seen = set()
    sentences = [s.strip() for s in re.split(r"(?<=[.!?])\s+", text or "") if s.strip()]
    for sentence in sentences:
        for pattern in GRAPH_RELATION_PATTERNS:
            for match in re.finditer(pattern, sentence, flags=re.IGNORECASE):
                source = normalize_entity(match.group(1))
                relation = normalize_entity(match.group(2)).lower()
                target = normalize_entity(match.group(3))
                if not source or not target or source == target:
                    continue
                if relation == "is" or relation == "are":
                    relation = "is"
                key = (source.lower(), relation, target.lower())
                if key in seen:
                    continue
                seen.add(key)
                facts.append({"source": source, "relation": relation, "target": target, "evidence": sentence})
    return facts


def build_graph_index_for_text(source_name: str, text: str):
    facts = extract_graph_facts(text)
    GRAPH_INDEX[source_name] = facts
    if NEO4J_DRIVER:
        with NEO4J_DRIVER.session() as session:
            session.run(
                "MERGE (d:Document {name: $source}) "
                "WITH d "
                "UNWIND $facts AS fact "
                "MERGE (a:Entity {name: fact.source}) "
                "MERGE (b:Entity {name: fact.target}) "
                "MERGE (a)-[:RELATION {name: fact.relation, evidence: fact.evidence, source: $source}]->(b) "
                "MERGE (d)-[:HAS_ENTITY]->(a) "
                "MERGE (d)-[:HAS_ENTITY]->(b)",
                source=source_name,
                facts=[{"source": fact["source"], "target": fact["target"], "relation": fact["relation"], "evidence": fact["evidence"]} for fact in facts],
            )
    return facts


def _ocr_image(img) -> str:
    """OCR a PIL image with RapidOCR (pure pip, no Tesseract install)."""
    global _ocr_engine
    import numpy as np
    if _ocr_engine is None:
        from rapidocr_onnxruntime import RapidOCR
        _ocr_engine = RapidOCR()
    result, _ = _ocr_engine(np.array(img.convert("RGB")))
    return "\n".join(r[1] for r in result) if result else ""


def parse(name: str, data: bytes):
    """Return list of (page_number, text)."""
    ext = name.lower().rsplit(".", 1)[-1]
    if ext == "pdf":
        r = PdfReader(io.BytesIO(data))
        out, pdf = [], None
        for i, p in enumerate(r.pages):
            text = p.extract_text() or ""
            if OCR_ENABLED and len(text.strip()) < 30:  # scanned / image-only page
                import pypdfium2 as pdfium
                pdf = pdf or pdfium.PdfDocument(data)
                text = _ocr_image(pdf[i].render(scale=2).to_pil())
            out.append((i + 1, text))
        return out
    if ext in ("png", "jpg", "jpeg", "bmp", "tiff", "webp") and OCR_ENABLED:
        from PIL import Image
        return [(1, _ocr_image(Image.open(io.BytesIO(data))))]
    if ext == "docx":
        d = docx.Document(io.BytesIO(data))
        return [(1, "\n".join(p.text for p in d.paragraphs))]
    if ext in ("txt", "md", "csv", "log"):
        return [(1, data.decode("utf-8", errors="ignore"))]
    raise HTTPException(400, f"Unsupported file type: .{ext}")


def chunk(text: str):
    words = text.split()
    step = max(1, CHUNK_WORDS - CHUNK_OVERLAP)
    for i in range(0, len(words), step):
        piece = " ".join(words[i:i + CHUNK_WORDS])
        if len(piece) > 40:
            yield piece
        if i + CHUNK_WORDS >= len(words):
            break


def audit(event: dict):
    event["ts"] = time.strftime("%Y-%m-%dT%H:%M:%S")
    with open(AUDIT_LOG, "a", encoding="utf-8") as f:
        f.write(json.dumps(event) + "\n")


def invoke_groq(context: str, question: str):
    candidates = []
    seen = set()
    for model_name in [GROQ_MODEL] + [m for m in GROQ_MODEL_FALLBACKS if m != GROQ_MODEL]:
        if model_name and model_name not in seen:
            candidates.append(model_name)
            seen.add(model_name)

    last_error = None
    for model_name in candidates:
        try:
            result = groq.chat.completions.create(
                model=model_name,
                temperature=0.1,
                messages=[
                    {"role": "system", "content": SYSTEM_PROMPT},
                    {"role": "user", "content": f"CONTEXT:\n{context}\n\nQUESTION: {question}"},
                ],
            )
            return result, model_name
        except Exception as exc:
            last_error = exc
            msg = str(exc).lower()
            if "model_not_found" in msg or "decommissioned" in msg or "does not exist" in msg:
                continue
            raise HTTPException(502, f"Groq error: {exc}")

    raise HTTPException(502, f"Groq error: {last_error}")


app = FastAPI(title="RAG Document Assistant")


class Login(BaseModel):
    username: str
    password: str


@app.post("/api/login")
def login(l: Login):
    now = time.time()
    recent = [t for t in FAILS.get(l.username, []) if now - t < 300]
    if len(recent) >= 5:
        raise HTTPException(429, "Too many failed attempts. Try again in 5 minutes.")
    rec = USERS.get(l.username)
    if not rec or not hmac.compare_digest(rec[0].encode(), l.password.encode()):
        FAILS[l.username] = recent + [now]
        audit({"event": "login_failed", "user": l.username})
        raise HTTPException(401, "Wrong username or password.")
    FAILS.pop(l.username, None)
    audit({"event": "login", "user": l.username})
    return {"token": signer.dumps({"u": l.username, "r": rec[1]}), "user": l.username, "role": rec[1]}


def current_user(authorization: str = Header(None)):
    if not authorization or not authorization.startswith("Bearer "):
        raise HTTPException(401, "Login required.")
    try:
        return signer.loads(authorization[7:], max_age=TOKEN_HOURS * 3600)
    except (BadSignature, SignatureExpired):
        raise HTTPException(401, "Session expired. Log in again.")


def admin_only(user=Depends(current_user)):
    if user["r"] != "admin":
        raise HTTPException(403, "Admin only.")
    return user


class Query(BaseModel):
    question: str
    filters: Optional[Dict[str, str]] = None
    top_k: Optional[int] = None


@app.post("/api/upload")
async def upload(file: UploadFile = File(...), user=Depends(admin_only)):
    data = await file.read()
    pages = parse(file.filename, data)
    col.delete(where={"source": file.filename})  # re-upload replaces old vectors
    ids, docs, metas = [], [], []
    combined_text = []
    for page, text in pages:
        combined_text.append(text)
        for c in chunk(text):
            ids.append(uuid.uuid4().hex)
            docs.append(c)
            metas.append({"source": file.filename, "page": page})
    if not docs:
        raise HTTPException(422, "No extractable text (scanned PDF? OCR is needed).")
    for i in range(0, len(docs), 100):
        col.add(ids=ids[i:i + 100], documents=docs[i:i + 100], metadatas=metas[i:i + 100])
    build_graph_index_for_text(file.filename, "\n".join(combined_text))
    audit({"event": "upload", "user": user["u"], "file": file.filename, "chunks": len(docs)})
    return {"file": file.filename, "chunks": len(docs), "status": "indexed"}


@app.post("/api/query")
def query(q: Query, user=Depends(current_user)):
    if col.count() == 0:
        raise HTTPException(400, "No documents indexed yet. Upload one first.")
    t0 = time.time()
    res = col.query(query_texts=[q.question], n_results=q.top_k or TOP_K, where=q.filters or None)
    docs, metas, dists = res["documents"][0], res["metadatas"][0], res["distances"][0]
    graph_facts = []
    for item in metas:
        source_name = item.get("source")
        graph_facts.extend(GRAPH_INDEX.get(source_name, []))
    graph_context = ""
    if graph_facts:
        graph_context = "\nGRAPH KNOWLEDGE:\n" + "\n".join(
            f"- {fact['source']} {fact['relation']} {fact['target']} | evidence: {fact['evidence']}"
            for fact in graph_facts[:12]
        )
    context = "\n\n".join(f"[Source: {m['source']}, page {m['page']}]\n{d}" for d, m in zip(docs, metas))
    if graph_context:
        context = f"{context}\n\n{graph_context}"
    out, used_model = invoke_groq(context, q.question)
    sources = [
        {"source": m["source"], "page": m["page"], "score": round(1 - d, 3), "excerpt": doc[:300]}
        for doc, m, d in zip(docs, metas, dists)
    ]
    graph_summary = []
    seen_graph = set()
    for fact in graph_facts:
        key = (fact["source"].lower(), fact["relation"], fact["target"].lower())
        if key in seen_graph:
            continue
        seen_graph.add(key)
        graph_summary.append({
            "source": fact["source"],
            "relation": fact["relation"],
            "target": fact["target"],
            "evidence": fact["evidence"][:140]
        })
    ms = int((time.time() - t0) * 1000)
    audit({"event": "query", "user": user["u"], "question": q.question, "retrieved": [(s["source"], s["page"]) for s in sources], "ms": ms})
    return {"answer": out.choices[0].message.content, "sources": sources, "graph": graph_summary[:8], "latency_ms": ms}


@app.get("/api/graph")
def graph(user=Depends(current_user)):
    if NEO4J_DRIVER:
        with NEO4J_DRIVER.session() as session:
            rows = session.run(
                "MATCH (n)-[r]->(m) RETURN n.name AS source, type(r) AS relation, m.name AS target, r.evidence AS evidence LIMIT 50"
            )
            nodes = []
            edges = []
            seen = set()
            for row in rows:
                src = row["source"]
                tgt = row["target"]
                for node in [src, tgt]:
                    if node and node.lower() not in seen:
                        seen.add(node.lower())
                        nodes.append({"id": node, "label": node})
                edges.append({
                    "from": src,
                    "to": tgt,
                    "label": row["relation"],
                    "evidence": row["evidence"],
                })
            return {"nodes": nodes, "edges": edges, "documents": list(GRAPH_INDEX.keys()), "source": "neo4j", "status": "live"}

    nodes = []
    edges = []
    seen = set()
    for source_name, facts in GRAPH_INDEX.items():
        for fact in facts:
            for node in [fact["source"], fact["target"]]:
                if node.lower() not in seen:
                    seen.add(node.lower())
                    nodes.append({"id": node, "label": node})
            edges.append({
                "from": fact["source"],
                "to": fact["target"],
                "label": fact["relation"],
                "source_file": source_name,
            })
    return {"nodes": nodes, "edges": edges, "documents": list(GRAPH_INDEX.keys()), "source": "local", "status": "fallback"}


@app.get("/api/sources")
def sources(user=Depends(current_user)):
    counts: Dict[str, int] = {}
    for m in col.get(include=["metadatas"])["metadatas"]:
        counts[m["source"]] = counts.get(m["source"], 0) + 1
    return [{"source": k, "chunks": v} for k, v in sorted(counts.items())]


@app.delete("/api/sources/{name}")
def delete_source(name: str, user=Depends(admin_only)):
    col.delete(where={"source": name})
    audit({"event": "delete", "file": name})
    return {"deleted": name}


@app.get("/api/logs")
def logs(n: int = 50, user=Depends(admin_only)):
    if not AUDIT_LOG.exists():
        return []
    return [json.loads(l) for l in AUDIT_LOG.read_text(encoding="utf-8").splitlines()[-n:]]


app.mount("/static", StaticFiles(directory=BASE / "static"), name="static")


@app.get("/")
def index():
    return FileResponse(BASE / "static" / "index.html")


if __name__ == "__main__":
    import uvicorn

    uvicorn.run("app:app", host="127.0.0.1", port=8000, reload=False)
